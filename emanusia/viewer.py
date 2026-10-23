"""Visualizador de arquivos dentro do sistema (sem abrir Excel, leitor de PDF etc.).
Planilhas viram uma grade com a formatação (R$, %, datas, cores, mesclagens) e os gráficos do
próprio arquivo; fórmulas são calculadas pelo motor do sistema (xlformula)."""
import csv
import datetime as dt
import io
import posixpath
import re
import zipfile
from pathlib import Path

from . import config

MAX_ROWS, MAX_COLS = 500, 60
TEXT_EXT = {".txt", ".log", ".json", ".xml", ".yaml", ".yml", ".md", ".vbs", ".sql", ".ini"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg"}


def view(path: str) -> dict:
    p = config.safe_path(path, mkdir=False)
    if not p.is_file():
        raise FileNotFoundError(path)
    ext = p.suffix.lower()
    base = {"path": config.display_path(p), "name": p.name, "size": p.stat().st_size, "ext": ext.lstrip(".")}
    if ext in (".xlsx", ".xlsm"):
        return {**base, "kind": "sheet", "sheets": _xlsx(p)}
    if ext in (".csv", ".tsv"):
        return {**base, "kind": "sheet", "sheets": [_csv(p)]}
    if ext in (".html", ".htm"):
        return {**base, "kind": "html"}
    if ext == ".pdf":
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(p))
        try:
            return {**base, "kind": "pdf", "pages": len(doc)}
        finally:
            doc.close()
    if ext in IMG_EXT:
        return {**base, "kind": "image"}
    if ext in TEXT_EXT:
        raw = p.read_bytes()[:400_000]
        for enc in ("utf-8-sig", "latin-1"):
            try:
                return {**base, "kind": "text", "text": raw.decode(enc)}
            except UnicodeDecodeError:
                continue
    return {**base, "kind": "other"}


# ------------------------------------------------------------------ CSV
def _csv(p: Path) -> dict:
    raw = p.read_bytes()[:3_000_000]
    for enc in ("utf-8-sig", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    total = len(rows)
    rows = [r[:MAX_COLS] for r in rows[:MAX_ROWS]]
    ncol = max((len(r) for r in rows), default=0)
    cells = [[{"v": v, **({"b": 1} if i == 0 else {})} for v in r] + [{"v": ""}] * (ncol - len(r)) for i, r in enumerate(rows)]
    return {"name": p.stem, "rows": cells, "widths": [110] * ncol, "heights": [], "merges": [], "charts": [],
            "truncated": total > MAX_ROWS}


# ------------------------------------------------------------------ XLSX
def _color(c) -> str | None:
    try:
        if c is not None and c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) in (6, 8):
            rgb = c.rgb[-6:]
            return None if c.rgb in ("00000000",) else "#" + rgb
    except Exception:
        return None
    return None


def fmt_value(v, number_format: str = "General") -> str:
    """Valor como o Excel mostraria (padrão brasileiro)."""
    from .xlformula import XLError
    if v is None:
        return ""
    if isinstance(v, XLError):
        return v.code
    if isinstance(v, bool):
        return "VERDADEIRO" if v else "FALSO"
    nf = (number_format or "General")
    if isinstance(v, (dt.datetime, dt.date)):
        has_time = isinstance(v, dt.datetime) and (v.hour or v.minute or "h" in nf.lower())
        return v.strftime("%d/%m/%Y %H:%M" if has_time else "%d/%m/%Y")
    if not isinstance(v, (int, float)):
        return str(v)
    br = lambda x, d: (f"{x:,.{d}f}").replace(",", "X").replace(".", ",").replace("X", ".")  # noqa: E731
    low = nf.lower()
    dec = len(m.group(1)) if (m := re.search(r"0\.(0+)", nf)) else 0
    if any(k in low for k in ("dd", "yy", "mm/")) and not isinstance(v, bool):
        try:
            d = dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))
            return d.strftime("%d/%m/%Y %H:%M" if "h" in low else "%d/%m/%Y")
        except (OverflowError, ValueError):
            pass
    if "%" in nf:
        return br(v * 100, dec) + "%"
    if "r$" in low or "$" in nf:
        return ("-" if v < 0 else "") + "R$ " + br(abs(v), dec if "0." in nf else 0)
    if "#,##0" in nf or "0." in nf or nf.strip() in ("0",):
        return br(v, dec)
    if isinstance(v, float) and not v.is_integer():
        return br(v, 2) if abs(v) >= 1 else br(v, 4).rstrip("0").rstrip(",")
    return br(v, 0) if abs(v) >= 1000 else str(int(v)) if float(v).is_integer() else str(v)


def _xlsx(p: Path) -> list[dict]:
    from openpyxl import load_workbook
    from .xlformula import NAME, Evaluator
    wb = load_workbook(p)
    try:
        cached = load_workbook(p, data_only=True)
    except Exception:
        cached = None
    ev = Evaluator(wb)
    charts = _charts(p, ev)
    out = []
    for ws in wb.worksheets:
        nrow, ncol = min(ws.max_row, MAX_ROWS), min(ws.max_column, MAX_COLS)
        rows = []
        for r in range(1, nrow + 1):
            line = []
            for c in range(1, ncol + 1):
                cell = ws.cell(r, c)
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    v = ev.value(ws.title, r, c)
                    if v == NAME and cached is not None:  # função que o motor não conhece: usa o valor salvo
                        v = cached[ws.title].cell(r, c).value
                d = {"v": fmt_value(v, cell.number_format)}
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    d["n"] = 1
                f = cell.font
                if f is not None:
                    if f.b:
                        d["b"] = 1
                    if f.sz and float(f.sz) != 11:
                        d["sz"] = float(f.sz)
                    if (col := _color(f.color)):
                        d["c"] = col
                if cell.fill is not None and cell.fill.fill_type == "solid" and (bg := _color(cell.fill.fgColor)):
                    d["bg"] = bg
                if cell.alignment is not None and cell.alignment.horizontal in ("center", "right", "left"):
                    d["al"] = cell.alignment.horizontal
                line.append(d)
            rows.append(line)
        from openpyxl.utils import get_column_letter
        widths = []
        for c in range(1, ncol + 1):
            w = ws.column_dimensions[get_column_letter(c)].width if get_column_letter(c) in ws.column_dimensions else None
            widths.append(int((w or 8.43) * 7 + 5))
        heights = [int((ws.row_dimensions[r].height or 15) * 4 / 3) for r in range(1, nrow + 1)]
        merges = [[m.min_row, m.min_col, m.max_row, m.max_col] for m in ws.merged_cells.ranges
                  if m.min_row <= nrow and m.min_col <= ncol]
        out.append({"name": ws.title, "rows": rows, "widths": widths, "heights": heights, "merges": merges,
                    "charts": charts.get(ws.title, []), "grid": ws.sheet_view.showGridLines is not False,
                    "truncated": ws.max_row > MAX_ROWS or ws.max_column > MAX_COLS})
    return out


# ------------------------------------------------------------------ gráficos do arquivo
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "pr": "http://schemas.openxmlformats.org/package/2006/relationships"}
EMU = 9525  # EMU por pixel


def _rels(z, part) -> dict:
    d, f = posixpath.split(part)
    name = f"{d}/_rels/{f}.rels"
    if name not in z.namelist():
        return {}
    from lxml import etree
    root = etree.fromstring(z.read(name))
    return {r.get("Id"): posixpath.normpath(posixpath.join(d, r.get("Target"))) if not r.get("Target").startswith("/")
            else r.get("Target").lstrip("/") for r in root.findall("pr:Relationship", NS)}


def _charts(p: Path, ev) -> dict:
    from lxml import etree
    out: dict = {}
    try:
        z = zipfile.ZipFile(p)
        wbx = etree.fromstring(z.read("xl/workbook.xml"))
        wrels = _rels(z, "xl/workbook.xml")
        for s in wbx.find("m:sheets", NS):
            sheet = s.get("name")
            part = wrels.get(s.get(f"{{{NS['r']}}}id"))
            if not part or part not in z.namelist():
                continue
            srels = _rels(z, part)
            sx = etree.fromstring(z.read(part))
            for dr in sx.findall("m:drawing", NS):
                dpart = srels.get(dr.get(f"{{{NS['r']}}}id"))
                if not dpart or dpart not in z.namelist():
                    continue
                drels = _rels(z, dpart)
                dx = etree.fromstring(z.read(dpart))
                for anc in dx:
                    ch = anc.find(".//c:chart", NS)
                    if ch is None:
                        continue
                    cpart = drels.get(ch.get(f"{{{NS['r']}}}id"))
                    if not cpart or cpart not in z.namelist():
                        continue
                    spec = _chart(etree.fromstring(z.read(cpart)), ev, sheet)
                    if spec:
                        spec["anchor"] = _anchor(anc)
                        out.setdefault(sheet, []).append(spec)
    except Exception:  # gráfico ilegível não impede ver a planilha
        import logging
        logging.getLogger("emanusia").warning("não consegui ler os gráficos de %s", p, exc_info=True)
    return out


def _anchor(anc) -> dict:
    def pos(tag):
        e = anc.find(f"xdr:{tag}", NS)
        if e is None:
            return None
        g = lambda k: int(e.findtext(f"xdr:{k}", "0", NS))  # noqa: E731
        return {"col": g("col"), "row": g("row"), "colOff": g("colOff") // EMU, "rowOff": g("rowOff") // EMU}
    a = {"from": pos("from"), "to": pos("to")}
    ext = anc.find("xdr:ext", NS)
    if ext is not None:
        a["w"], a["h"] = int(ext.get("cx")) // EMU, int(ext.get("cy")) // EMU
    return a


def _vals(node, ev, sheet, numeric):
    """Valores de c:tx / c:cat / c:val: pela referência (calculada) ou pelo cache gravado."""
    from .xlformula import Range
    if node is None:
        return []
    f = node.findtext(".//c:f", None, NS)
    if f:
        r = ev.ref(f, sheet)
        vals = r.flat() if isinstance(r, Range) else [r]
    else:
        vals = [pt.findtext("c:v", "", NS) for pt in node.iterfind(".//c:pt", NS)]
    if numeric:
        out = []
        for v in vals:
            try:
                out.append(float(v) if v not in (None, "") else None)
            except (TypeError, ValueError):
                out.append(None)
        return out
    return [fmt_value(v) if not isinstance(v, str) else v for v in vals]


def _chart(cx, ev, sheet) -> dict | None:
    title = "".join(cx.find("c:chart/c:title", NS).itertext()).strip() if cx.find("c:chart/c:title", NS) is not None else ""
    title = "" if title == "None" else title
    plot = cx.find("c:chart/c:plotArea", NS)
    if plot is None:
        return None
    series, kind = [], None
    for g in plot:
        tag = etree_local(g.tag)
        if not tag.endswith("Chart"):
            continue
        k = {"barChart": "bar", "bar3DChart": "bar", "lineChart": "line", "line3DChart": "line", "pieChart": "pie",
             "pie3DChart": "pie", "doughnutChart": "doughnut", "areaChart": "area", "area3DChart": "area",
             "scatterChart": "scatter"}.get(tag)
        if not k:
            continue
        kind = kind or k
        horizontal = g.find("c:barDir", NS) is not None and g.find("c:barDir", NS).get("val") == "bar"
        for s in g.findall("c:ser", NS):
            name = _vals(s.find("c:tx", NS), ev, sheet, False)
            cats = _vals(s.find("c:cat", NS) if s.find("c:cat", NS) is not None else s.find("c:xVal", NS), ev, sheet, False)
            vals = _vals(s.find("c:val", NS) if s.find("c:val", NS) is not None else s.find("c:yVal", NS), ev, sheet, True)
            n = len(vals)
            series.append({"type": k, "name": (name or [""])[0], "x": cats[:n] if cats else list(range(1, n + 1)),
                           "y": vals, "horizontal": horizontal})
    if not series:
        return None
    return {"title": title, "type": kind, "series": series}


def etree_local(tag) -> str:
    return tag.split("}", 1)[1] if isinstance(tag, str) and "}" in tag else str(tag)


def pdf_page_png(path: str, page: int, scale: float = 1.6) -> bytes:
    """Página do PDF como imagem, desenhada pelo próprio sistema (não depende do leitor de PDF do navegador)."""
    import pypdfium2 as pdfium
    p = config.safe_path(path, mkdir=False)
    doc = pdfium.PdfDocument(str(p))
    try:
        if not 1 <= page <= len(doc):
            raise FileNotFoundError(f"página {page}")
        img = doc[page - 1].render(scale=max(0.5, min(float(scale), 3.0))).to_pil()
        buf = io.BytesIO()
        img.save(buf, "PNG", optimize=True)
        return buf.getvalue()
    finally:
        doc.close()
