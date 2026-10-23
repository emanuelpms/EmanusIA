"""Motor de fórmulas do Excel, sem o Excel: calcula as fórmulas de um .xlsx (SOMA, SOMASE, CONT.SE,
MÉDIASE, MÁXIMOSES, SE, SEERRO, PROCV, ARRED...) e grava os valores calculados dentro do arquivo.
Assim a planilha/dashboard mostra os números em qualquer lugar — no visualizador do sistema, no
celular, no Excel Online — e os próximos passos do robô conseguem ler os resultados.
As fórmulas continuam no arquivo: quem abrir no Excel vê tudo recalcular normalmente."""
import datetime as dt
import math
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.utils import column_index_from_string, get_column_letter


class XLError:
    def __init__(self, code): self.code = code
    def __repr__(self): return self.code
    def __eq__(self, o): return isinstance(o, XLError) and o.code == self.code
    def __hash__(self): return hash(self.code)


DIV0, VALUE, REF, NAME, NA = (XLError(c) for c in ("#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#N/A"))


class Range:
    """Bloco de células já avaliadas (linhas x colunas)."""
    def __init__(self, rows): self.rows = rows
    def flat(self): return [v for r in self.rows for v in r]


_REF = re.compile(r"^(?:(?P<sheet>'(?:[^']|'')+'|[^!]+)!)?(?P<a>\$?[A-Z]{1,3}\$?\d*|\$?\d+)(?::(?P<b>\$?[A-Z]{1,3}\$?\d*|\$?\d+))?$", re.I)


def _split(ref):
    ref = ref.replace("$", "")
    m = re.match(r"^([A-Z]{0,3})(\d*)$", ref, re.I)
    col = column_index_from_string(m.group(1).upper()) if m.group(1) else None
    row = int(m.group(2)) if m.group(2) else None
    return col, row


class Evaluator:
    def __init__(self, wb):
        self.wb = wb
        self.cache: dict = {}
        self.busy: set = set()

    # ------------------------------------------------ células
    def value(self, sheet: str, row: int, col: int):
        key = (sheet, row, col)
        if key in self.cache:
            return self.cache[key]
        try:
            ws = self.wb[sheet]
        except KeyError:
            return REF
        v = ws.cell(row, col).value
        if isinstance(v, str) and v.startswith("="):
            if key in self.busy:
                return REF  # referência circular
            self.busy.add(key)
            try:
                v = self.formula(v, sheet)
            finally:
                self.busy.discard(key)
        elif hasattr(v, "text") and not isinstance(v, str):  # fórmulas de matriz/objetos: não suportado
            v = NAME
        self.cache[key] = v
        return v

    def ref(self, text: str, sheet: str):
        m = _REF.match(text.strip())
        if not m:
            return NAME
        sh = m.group("sheet")
        sh = sh.strip("'").replace("''", "'") if sh else sheet
        if sh not in self.wb.sheetnames:
            return REF
        ws = self.wb[sh]
        c1, r1 = _split(m.group("a"))
        if not m.group("b"):
            return self.value(sh, r1, c1) if c1 and r1 else REF
        c2, r2 = _split(m.group("b"))
        c1, c2 = c1 or 1, c2 or ws.max_column
        r1, r2 = r1 or 1, min(r2 or ws.max_row, ws.max_row) if not r2 or r2 > ws.max_row else r2
        if m.group("a").replace("$", "").isalpha():  # coluna inteira (A:A): só até a última linha usada
            r1, r2 = 1, ws.max_row
        return Range([[self.value(sh, r, c) for c in range(min(c1, c2), max(c1, c2) + 1)]
                      for r in range(min(r1, r2), max(r1, r2) + 1)])

    # ------------------------------------------------ fórmula
    def formula(self, text: str, sheet: str):
        try:
            toks = [t for t in Tokenizer(text).items if t.type != Token.WSPACE]
        except Exception:
            return NAME
        p = _Parser(toks, self, sheet)
        try:
            v = p.expr()
            if p.i != len(toks):
                return NAME
        except _Err as e:
            return e.err
        except (ZeroDivisionError,):
            return DIV0
        except (TypeError, ValueError, IndexError, KeyError, OverflowError):
            return VALUE
        if isinstance(v, Range):
            v = v.rows[0][0] if v.rows and v.rows[0] else 0
        return 0 if v is None else v


class _Err(Exception):
    def __init__(self, err): self.err = err


def _num(v):
    if isinstance(v, XLError):
        raise _Err(v)
    if v is None or v == "":
        return 0
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, (dt.datetime, dt.date)):
        return _serial(v)
    try:
        return float(str(v).replace(",", ".")) if re.fullmatch(r"\s*-?\d+([.,]\d+)?\s*", str(v)) else _raise(VALUE)
    except _Err:
        raise


def _raise(e):
    raise _Err(e)


def _serial(d):
    if isinstance(d, dt.datetime):
        return (d - dt.datetime(1899, 12, 30)).total_seconds() / 86400
    return (d - dt.date(1899, 12, 30)).days


def _txt(v):
    if isinstance(v, XLError):
        raise _Err(v)
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VERDADEIRO" if v else "FALSO"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _scalar(v):
    return v.rows[0][0] if isinstance(v, Range) and v.rows and v.rows[0] else v


def _nums(args):
    """Números dos argumentos (em intervalos, só células numéricas — como o Excel)."""
    out = []
    for a in args:
        if isinstance(a, Range):
            for v in a.flat():
                if isinstance(v, XLError):
                    raise _Err(v)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    out.append(v)
                elif isinstance(v, (dt.datetime, dt.date)):
                    out.append(_serial(v))
        else:
            out.append(_num(a))
    return out


def _crit(c):
    """Critério do SOMASE/CONT.SE -> função(valor) -> bool."""
    c = _scalar(c)
    if isinstance(c, (int, float)) and not isinstance(c, bool):
        return lambda v: _eqnum(v, c)
    s = _txt(c)
    m = re.match(r"^(<=|>=|<>|<|>|=)?(.*)$", s, re.S)
    op, rest = m.group(1) or "=", m.group(2)
    try:
        n = float(rest.replace(",", ".")) if rest.strip() else None
    except ValueError:
        n = None
    if op in ("<", ">", "<=", ">="):
        f = {"<": lambda a, b: a < b, ">": lambda a, b: a > b, "<=": lambda a, b: a <= b, ">=": lambda a, b: a >= b}[op]
        if n is not None:
            return lambda v: isinstance(v, (int, float)) and not isinstance(v, bool) and f(v, n)
        return lambda v: isinstance(v, str) and f(v.lower(), rest.lower())
    if n is not None:
        test = lambda v: _eqnum(v, n) or (isinstance(v, str) and v.strip().lower() == rest.lower())  # noqa: E731
    elif rest == "":
        test = lambda v: v is None or v == ""  # noqa: E731
    else:
        rx = re.compile("^" + re.escape(rest.lower()).replace(r"\*", ".*").replace(r"\?", ".") + "$", re.S)
        test = lambda v: v is not None and not isinstance(v, (int, float)) and bool(rx.match(_txt(v).lower()))  # noqa: E731
    return (lambda v: not test(v)) if op == "<>" else test


def _eqnum(v, n):
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return abs(v - n) < 1e-9
    if isinstance(v, str):
        try:
            return abs(float(v.replace(",", ".")) - n) < 1e-9
        except ValueError:
            return False
    return False


def _ifs(values, pairs):
    """Linhas que atendem a todos os critérios (pares intervalo, critério)."""
    vals = values.flat() if values is not None else None
    tests = [(r.flat(), _crit(c)) for r, c in pairs]
    n = len(tests[0][0]) if tests else 0
    for flat, _ in tests:
        if len(flat) != n:
            raise _Err(VALUE)
    idx = [i for i in range(n) if all(t(f[i]) for f, t in tests)]
    return idx, vals


def _agg_if(kind):
    def fn(*a):
        if kind in ("SUMIF", "AVERAGEIF"):
            rng, crit = a[0], a[1]
            vals = a[2] if len(a) > 2 else rng
            idx, flat = _ifs(vals, [(rng, crit)])
        elif kind == "COUNTIF":
            idx, flat = _ifs(None, [(a[0], a[1])])
            return len(idx)
        elif kind == "COUNTIFS":
            idx, flat = _ifs(None, list(zip(a[0::2], a[1::2])))
            return len(idx)
        else:  # SUMIFS AVERAGEIFS MAXIFS MINIFS: (valores, intervalo1, critério1, ...)
            idx, flat = _ifs(a[0], list(zip(a[1::2], a[2::2])))
        nums = [flat[i] for i in idx if isinstance(flat[i], (int, float)) and not isinstance(flat[i], bool)]
        if kind in ("SUMIF", "SUMIFS"):
            return sum(nums)
        if kind in ("AVERAGEIF", "AVERAGEIFS"):
            return sum(nums) / len(nums) if nums else _raise(DIV0)
        if kind == "MAXIFS":
            return max(nums) if nums else 0
        return min(nums) if nums else 0
    return fn


def _vlookup(key, table, col, approx=True):
    key, col = _scalar(key), int(_num(_scalar(col)))
    rows = table.rows
    if col < 1 or (rows and col > len(rows[0])):
        raise _Err(REF)
    exact = not (approx if not isinstance(approx, Range) else _scalar(approx))
    if exact:
        k = _txt(key).lower() if not isinstance(key, (int, float)) else key
        for r in rows:
            v = r[0]
            if (isinstance(k, (int, float)) and _eqnum(v, k)) or (not isinstance(k, (int, float)) and _txt(v).lower() == k):
                return r[col - 1]
        raise _Err(NA)
    best = None
    for r in rows:
        try:
            if r[0] is not None and r[0] <= key:
                best = r
        except TypeError:
            continue
    if best is None:
        raise _Err(NA)
    return best[col - 1]


def _round(v, d=0, mode="half"):
    v, d = _num(_scalar(v)), int(_num(_scalar(d)))
    f = 10 ** d
    if mode == "up":
        return math.copysign(math.ceil(abs(v) * f) / f, v)
    if mode == "down":
        return math.copysign(math.floor(abs(v) * f) / f, v)
    return math.copysign(math.floor(abs(v) * f + 0.5) / f, v)


def _count(args, numeric=True):
    n = 0
    for a in args:
        for v in (a.flat() if isinstance(a, Range) else [a]):
            if numeric:
                n += isinstance(v, (int, float, dt.date)) and not isinstance(v, bool)
            else:
                n += v is not None and v != ""
    return n


def _bool(v):
    v = _scalar(v)
    if isinstance(v, XLError):
        raise _Err(v)
    if isinstance(v, str):
        if v.upper() in ("TRUE", "VERDADEIRO"):
            return True
        if v.upper() in ("FALSE", "FALSO"):
            return False
        raise _Err(VALUE)
    return bool(v)


FUNCS = {
    "SUM": lambda *a: sum(_nums(a)),
    "AVERAGE": lambda *a: (lambda n: sum(n) / len(n) if n else _raise(DIV0))(_nums(a)),
    "MAX": lambda *a: max(_nums(a) or [0]), "MIN": lambda *a: min(_nums(a) or [0]),
    "COUNT": lambda *a: _count(a), "COUNTA": lambda *a: _count(a, False),
    "COUNTBLANK": lambda r: sum(1 for v in r.flat() if v is None or v == ""),
    "SUMIF": _agg_if("SUMIF"), "SUMIFS": _agg_if("SUMIFS"), "COUNTIF": _agg_if("COUNTIF"),
    "COUNTIFS": _agg_if("COUNTIFS"), "AVERAGEIF": _agg_if("AVERAGEIF"), "AVERAGEIFS": _agg_if("AVERAGEIFS"),
    "MAXIFS": _agg_if("MAXIFS"), "MINIFS": _agg_if("MINIFS"),
    "ROUND": lambda v, d=0: _round(v, d), "ROUNDUP": lambda v, d=0: _round(v, d, "up"),
    "ROUNDDOWN": lambda v, d=0: _round(v, d, "down"), "INT": lambda v: math.floor(_num(_scalar(v))),
    "ABS": lambda v: abs(_num(_scalar(v))),
    "AND": lambda *a: all(_bool(x) for x in a), "OR": lambda *a: any(_bool(x) for x in a),
    "NOT": lambda v: not _bool(v), "TRUE": lambda: True, "FALSE": lambda: False,
    "CONCATENATE": lambda *a: "".join(_txt(_scalar(x)) for x in a), "CONCAT": lambda *a: "".join(_txt(_scalar(x)) for x in a),
    "LEN": lambda v: len(_txt(_scalar(v))), "UPPER": lambda v: _txt(_scalar(v)).upper(),
    "LOWER": lambda v: _txt(_scalar(v)).lower(), "TRIM": lambda v: " ".join(_txt(_scalar(v)).split()),
    "LEFT": lambda v, n=1: _txt(_scalar(v))[:int(_num(_scalar(n)))],
    "RIGHT": lambda v, n=1: _txt(_scalar(v))[-int(_num(_scalar(n))):] if int(_num(_scalar(n))) else "",
    "TODAY": lambda: dt.date.today(), "NOW": lambda: dt.datetime.now().replace(microsecond=0),
    "VLOOKUP": _vlookup, "IFNA": None, "IF": None, "IFERROR": None,  # IF/IFERROR/IFNA: tratados à parte
    "YEAR": lambda v: _date(v).year, "MONTH": lambda v: _date(v).month, "DAY": lambda v: _date(v).day,
}


def _date(v):
    v = _scalar(v)
    if isinstance(v, (dt.date, dt.datetime)):
        return v
    return dt.datetime(1899, 12, 30) + dt.timedelta(days=_num(v))


class _Parser:
    def __init__(self, toks, ev: Evaluator, sheet):
        self.t, self.i, self.ev, self.sheet = toks, 0, ev, sheet

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else None

    def take(self):
        tok = self.t[self.i]
        self.i += 1
        return tok

    def _op(self, values):
        p = self.peek()
        return p is not None and p.type == Token.OP_IN and p.value in values

    def expr(self):
        a = self.concat()
        while self._op(("=", "<>", "<", ">", "<=", ">=")):
            op = self.take().value
            b = self.concat()
            a = _compare(op, _scalar(a), _scalar(b))
        return a

    def concat(self):
        a = self.add()
        while self._op(("&",)):
            self.take()
            a = _txt(_scalar(a)) + _txt(_scalar(self.add()))
        return a

    def add(self):
        a = self.mul()
        while self._op(("+", "-")):
            op = self.take().value
            b = self.mul()
            a = _num(_scalar(a)) + _num(_scalar(b)) if op == "+" else _num(_scalar(a)) - _num(_scalar(b))
        return a

    def mul(self):
        a = self.pow()
        while self._op(("*", "/")):
            op = self.take().value
            b = self.pow()
            if op == "*":
                a = _num(_scalar(a)) * _num(_scalar(b))
            else:
                d = _num(_scalar(b))
                a = _raise(DIV0) if d == 0 else _num(_scalar(a)) / d
        return a

    def pow(self):
        a = self.unary()
        while self._op(("^",)):
            self.take()
            a = _num(_scalar(a)) ** _num(_scalar(self.unary()))
        return a

    def unary(self):
        p = self.peek()
        if p is not None and p.type == Token.OP_PRE:
            self.take()
            v = self.unary()
            return -_num(_scalar(v)) if p.value == "-" else _num(_scalar(v))
        v = self.primary()
        p = self.peek()
        if p is not None and p.type == Token.OP_POST and p.value == "%":
            self.take()
            v = _num(_scalar(v)) / 100
        return v

    def primary(self):
        tok = self.take()
        if tok.type == Token.OPERAND:
            if tok.subtype == Token.NUMBER:
                return float(tok.value) if any(c in tok.value for c in ".eE") else int(tok.value)
            if tok.subtype == Token.TEXT:
                return tok.value[1:-1].replace('""', '"')
            if tok.subtype == Token.LOGICAL:
                return tok.value.upper() == "TRUE"
            if tok.subtype == Token.ERROR:
                return XLError(tok.value)
            return self.ev.ref(tok.value, self.sheet)
        if tok.type == Token.PAREN and tok.subtype == Token.OPEN:
            v = self.expr()
            self.take()  # )
            return v
        if tok.type == Token.FUNC and tok.subtype == Token.OPEN:
            name = tok.value[:-1].upper().removeprefix("_XLFN.").removeprefix("_XLWS.")
            args = []
            if not (self.peek().type == Token.FUNC and self.peek().subtype == Token.CLOSE):
                while True:
                    p = self.peek()
                    if p.type == Token.SEP or (p.type == Token.FUNC and p.subtype == Token.CLOSE):
                        args.append(None)  # argumento vazio
                    else:
                        args.append(self._arg(name))
                    p = self.take()
                    if p.type == Token.FUNC and p.subtype == Token.CLOSE:
                        break
            else:
                self.take()
            return self._call(name, args)
        raise _Err(NAME)

    def _arg(self, name):
        """Argumento como valor (erros viram valor, para o SEERRO/SE conseguirem tratar)."""
        try:
            return self.expr()
        except _Err as e:
            return e.err
        except ZeroDivisionError:
            return DIV0

    def _call(self, name, args):
        if name in ("IFERROR", "IFNA"):
            v = _scalar(args[0])
            bad = isinstance(v, XLError) and (name == "IFERROR" or v == NA)
            return (args[1] if len(args) > 1 else 0) if bad else v
        if name == "IF":
            cond = _bool(args[0])
            return (args[1] if len(args) > 1 else True) if cond else (args[2] if len(args) > 2 else False)
        fn = FUNCS.get(name)
        if fn is None:
            raise _Err(NAME)
        for a in args:
            if isinstance(a, XLError):
                raise _Err(a)
        return fn(*[a for a in args if a is not None] if name not in ("VLOOKUP",) else args)


def _compare(op, a, b):
    for v in (a, b):
        if isinstance(v, XLError):
            raise _Err(v)
    if isinstance(a, str) and isinstance(b, str):
        a, b = a.lower(), b.lower()
    elif isinstance(a, str) != isinstance(b, str):
        if op in ("=", "<>"):
            return op == "<>"
        a, b = (1, 0) if isinstance(a, str) else (0, 1)  # no Excel texto > número
    a = 0 if a is None else a
    b = 0 if b is None else b
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]


# ------------------------------------------------------------------ arquivo inteiro
def evaluate(path) -> dict:
    """{aba: {(linha, coluna): valor}} só das células com fórmula."""
    wb = load_workbook(path)
    ev = Evaluator(wb)
    out = {}
    for ws in wb.worksheets:
        vals = {}
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    vals[(c.row, c.column)] = ev.value(ws.title, c.row, c.column)
        out[ws.title] = vals
    return out


def fill_cached_values(path) -> int:
    """Grava os valores calculados (<v>) ao lado das fórmulas, dentro do .xlsx. Devolve quantas."""
    from lxml import etree
    path = Path(path)
    vals = evaluate(path)
    if not any(vals.values()):
        return 0
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
        infos = {i.filename: i for i in z.infolist()}
    wbx = etree.fromstring(files["xl/workbook.xml"])
    rels = etree.fromstring(files["xl/_rels/workbook.xml.rels"])
    target = {r.get("Id"): r.get("Target") for r in rels}
    n = 0
    for s in wbx.iter(f"{{{ns}}}sheet"):
        name = s.get("name")
        t = target.get(s.get(f"{{{rel_ns}}}id"), "")
        part = t.lstrip("/") if t.startswith("/") else "xl/" + t
        if part not in files or not vals.get(name):
            continue
        root = etree.fromstring(files[part])
        for c in root.iter(f"{{{ns}}}c"):
            f = c.find(f"{{{ns}}}f")
            if f is None:
                continue
            col, row = _split(c.get("r"))
            if (row, col) not in vals[name]:
                continue
            v = vals[name][(row, col)]
            old = c.find(f"{{{ns}}}v")
            if old is not None:
                c.remove(old)
            ve = etree.SubElement(c, f"{{{ns}}}v")
            if isinstance(v, XLError):
                c.set("t", "e"); ve.text = v.code
            elif isinstance(v, bool):
                c.set("t", "b"); ve.text = "1" if v else "0"
            elif isinstance(v, (int, float)):
                c.attrib.pop("t", None); ve.text = repr(float(v)) if isinstance(v, float) else str(v)
            elif isinstance(v, (dt.datetime, dt.date)):
                c.attrib.pop("t", None); ve.text = repr(float(_serial(v)))
            else:
                c.set("t", "str"); ve.text = _txt(v)
            n += 1
        files[part] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(infos[name], data)
    tmp.replace(path)
    return n


def letter(col: int) -> str:
    return get_column_letter(col)
