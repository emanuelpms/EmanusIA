"""Excel profissional (openpyxl): relatórios formatados, dashboards nativos do Excel com fórmulas vivas,
edição de planilhas existentes e (no Windows) atualização de consultas/tabelas dinâmicas e macros."""
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .. import config
from .registry import tool

NAVY, BLUE, LIGHT = "1F3A5F", "2F6DB5", "EEF3FA"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
FORMATS = {"money": 'R$ #,##0.00', "moeda": 'R$ #,##0.00', "int": "#,##0", "inteiro": "#,##0",
           "decimal": "#,##0.00", "percent": "0.0%", "percentual": "0.0%", "date": "dd/mm/yyyy",
           "data": "dd/mm/yyyy", "datetime": "dd/mm/yyyy hh:mm"}
PALETTE = ["2F6DB5", "22A38A", "F2A33A", "E0605E", "8A6FD1", "3BA9D9", "D96BA8", "7FB43A"]
THIN = Side(style="thin", color="D6DEEA")


def _fmt(name: str) -> str:
    return FORMATS.get(str(name).lower(), name)  # aceita também formato Excel direto


def _col_index(ws, col) -> int:
    if isinstance(col, int):
        return col
    for c in ws[1]:
        if str(c.value).strip().lower() == str(col).strip().lower():
            return c.column
    raise KeyError(f"Coluna '{col}' não existe na aba '{ws.title}'")


def _fill_sheet(ws, rows: list[list], as_table: bool = False):
    for r in rows:
        ws.append(list(r))
    if ws.max_row == 0:
        return
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    if as_table and ws.max_row > 1:
        name = "".join(ch for ch in ws.title if ch.isalnum()) or "Tabela"
        t = Table(displayName=f"T_{name}"[:250], ref=ws.dimensions)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    elif ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 10), 50)


def _chart(ws, spec: dict, max_row: int | None = None):
    kind = spec.get("type", "bar")
    ch = {"bar": BarChart, "col": BarChart, "line": LineChart, "pie": PieChart, "area": AreaChart}.get(kind, BarChart)()
    if kind == "bar" and spec.get("horizontal"):
        ch.type = "bar"
    ch.title = spec.get("title", "")
    ch.style = 10
    last = max_row or ws.max_row
    for col in spec.get("value_cols", [2]):
        ch.add_data(Reference(ws, min_col=_col_index(ws, col), min_row=1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=_col_index(ws, spec.get("category_col", 1)), min_row=2, max_row=last))
    if kind == "pie":
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showPercent = True
    else:
        for i, s in enumerate(ch.series):
            s.graphicalProperties.solidFill = PALETTE[i % len(PALETTE)]
            s.graphicalProperties.line.solidFill = PALETTE[i % len(PALETTE)]
        ch.legend.position = "b"
    ch.width, ch.height = spec.get("width", 17), spec.get("height", 8.5)
    return ch


@tool("excel_write",
      "Cria um .xlsx profissional. sheets={\"Aba\": [[cabeçalho...], [linha...]]}. Opcionais por aba: "
      "formats={Aba: {coluna: money|int|decimal|percent|date}}, totals={Aba: [colunas p/ somar]}, "
      "conditional={Aba: {coluna: scale|bar}}, as_table=true (tabela do Excel), "
      "charts=[{sheet, type: bar|line|pie|area, category_col, value_cols:[...], title}] (colunas por nome ou nº).",
      {"path": {"type": "string"}, "sheets": {"type": "object"}, "formats": {"type": "object"},
       "totals": {"type": "object"}, "conditional": {"type": "object"}, "as_table": {"type": "boolean"},
       "charts": {"type": "array"}, "chart": {"type": "object"}},
      ("path", "sheets"))
def excel_write(path: str, sheets: dict, formats: dict | None = None, totals: dict | None = None,
                conditional: dict | None = None, as_table: bool = False, charts: list | None = None,
                chart: dict | None = None):
    from .data import load, rows as to_rows
    wb = Workbook()
    wb.remove(wb.active)
    data_rows = {}
    for name, rws in sheets.items():
        if not (isinstance(rws, list) and (not rws or isinstance(rws[0], list | tuple))):
            rws = to_rows(load(rws), None)  # arquivo, resultado de outro passo ou lista de {coluna: valor}
        ws = wb.create_sheet(str(name)[:31])
        _fill_sheet(ws, rws, as_table and not (totals or {}).get(name))
        data_rows[ws.title] = ws.max_row
    for name, cols in (formats or {}).items():
        ws = wb[str(name)[:31]]
        for col, f in cols.items():
            ci = _col_index(ws, col)
            for (cell,) in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
                cell.number_format = _fmt(f)
    for name, cols in (totals or {}).items():
        ws = wb[str(name)[:31]]
        last, r = ws.max_row, ws.max_row + 1
        ws.cell(r, 1, "TOTAL").font = Font(bold=True)
        for col in cols:
            ci, L = _col_index(ws, col), get_column_letter(_col_index(ws, col))
            c = ws.cell(r, ci, f"=SUM({L}2:{L}{last})")
            c.font, c.number_format = Font(bold=True), ws.cell(2, ci).number_format
        for c in ws[r]:
            c.fill, c.border = PatternFill("solid", fgColor=LIGHT), Border(top=Side(style="medium", color=NAVY))
    for name, cols in (conditional or {}).items():
        ws = wb[str(name)[:31]]
        for col, kind in cols.items():
            L = get_column_letter(_col_index(ws, col))
            rng = f"{L}2:{L}{data_rows[ws.title]}"
            ws.conditional_formatting.add(rng, DataBarRule(start_type="min", end_type="max", color="5B9BD5")
                                          if kind == "bar" else ColorScaleRule(start_type="min", start_color="F8696B",
                                          mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                          end_type="max", end_color="63BE7B"))
    for i, spec in enumerate([*(charts or []), *([chart] if chart else [])]):
        ws = wb[str(spec.get("sheet") or list(sheets)[0])[:31]]
        anchor = spec.get("anchor") or f"{get_column_letter(ws.max_column + 2)}{2 + i * 18}"
        ws.add_chart(_chart(ws, spec, data_rows[ws.title]), anchor)
    wb.save(config.safe_path(path))
    return {"path": path, "sheets": list(wb.sheetnames)}


@tool("excel_save", "Monta uma planilha a partir de dados (de outro passo ou de um arquivo): uma aba, com "
      "formato por coluna (money|int|decimal|percent|date), linha de TOTAL, tabela do Excel e um gráfico opcional. "
      "append=true acrescenta as linhas no fim de uma planilha que já existe.",
      {"path": {"type": "string"}, "data": {}, "sheet": {"type": "string"}, "formats": {"type": "object"},
       "totals": {"type": "array"}, "as_table": {"type": "boolean"}, "chart_type": {"type": "string"},
       "chart_x": {"type": "string"}, "chart_y": {"type": "array"}, "chart_title": {"type": "string"},
       "append": {"type": "boolean"}}, ("path", "data"))
def excel_save(path: str, data, sheet: str = "Dados", formats: dict | None = None, totals: list | None = None,
               as_table: bool = True, chart_type: str = "", chart_x: str = "", chart_y: list | None = None,
               chart_title: str = "", append: bool = False):
    from .data import load, rows as to_rows
    table = data if isinstance(data, list) and data and isinstance(data[0], list | tuple) else to_rows(load(data), None)
    sheet = (sheet or "Dados")[:31]
    p = config.safe_path(path)
    if append and p.exists():
        wb = load_workbook(p)
        if sheet in wb.sheetnames:
            ws = wb[sheet]
            head = [str(c.value) for c in ws[1]]
            new_head = [str(c) for c in table[0]]
            idx = [new_head.index(h) if h in new_head else None for h in head]
            for r in table[1:]:  # casa as colunas pelo nome
                ws.append([r[i] if i is not None and i < len(r) else None for i in idx])
            wb.save(p)
            return {"path": path, "added": len(table) - 1, "rows": ws.max_row - 1}
    if isinstance(chart_y, str):
        chart_y = [chart_y]
    chart = {"sheet": sheet, "type": chart_type, "category_col": chart_x or 1, "value_cols": chart_y or [2],
             "title": chart_title} if chart_type else None
    res = excel_write(path, {sheet: table}, {sheet: formats} if formats else None,
                      {sheet: totals} if totals else None, None, as_table, None, chart)
    return {**res, "rows": len(table) - 1}


@tool("excel_dashboard",
      "Cria um DASHBOARD NATIVO do Excel a partir de uma base: aba Dashboard com cartões de KPI e gráficos, "
      "aba Dados (tabela) e resumos com fórmulas SOMASE/CONT.SE que se atualizam quando os dados mudam. "
      "kpis=[{label, column, agg: sum|mean|count|max|min|nunique, format}], "
      "charts=[{type: bar|line|pie|area, x (coluna categoria), y (coluna valor), agg, title, top, period: month}].",
      {"path": {"type": "string"}, "source": {}, "sheet": {"type": "string"}, "title": {"type": "string"},
       "subtitle": {"type": "string"}, "kpis": {"type": "array"}, "charts": {"type": "array"},
       "formats": {"type": "object"}}, ("path", "source", "title"))
def excel_dashboard(path: str, source, title: str, sheet: str | None = None, subtitle: str = "",
                    kpis: list | None = None, charts: list | None = None, formats: dict | None = None):
    import pandas as pd
    from .data import _agg, _col, load, rows as to_rows
    df = load(source, sheet)
    for i, ch in enumerate(charts or []):  # colunas auxiliares (mês) para agrupar datas
        if ch.get("period") == "month":
            x = _col(df, ch["x"])
            aux = f"{x} (mês)"
            df[aux] = pd.to_datetime(df[x], errors="coerce", dayfirst=True).dt.strftime("%Y-%m")
            ch = charts[i] = {**ch, "x": aux, "_sort": "x"}
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    data = wb.create_sheet("Dados")
    _fill_sheet(data, to_rows(df, None), as_table=True)
    for col, f in (formats or {}).items():
        ci = _col_index(data, col)
        for (cell,) in data.iter_rows(min_row=2, min_col=ci, max_col=ci):
            cell.number_format = _fmt(f)
    calc = wb.create_sheet("Resumos")
    letter = {str(c.value): get_column_letter(c.column) for c in data[1]}

    def rng(col):
        return f"Dados!${letter[col]}:${letter[col]}"

    # ---- layout do painel
    dash.sheet_view.showGridLines = False
    for i in range(1, 30):
        dash.column_dimensions[get_column_letter(i)].width = 11
    dash.merge_cells("B2:P2")
    dash["B2"] = title
    dash["B2"].font = Font(size=22, bold=True, color=NAVY)
    dash["B3"] = subtitle or f"{len(df)} registros · atualiza ao abrir no Excel"
    dash["B3"].font = Font(size=11, color="6B7A90")
    fn = {"sum": "SUM({r})", "mean": "AVERAGE({r})", "count": "COUNTA({r})-1", "max": "MAX({r})",
          "min": "MIN({r})"}
    for k, kp in enumerate(kpis or []):
        c0 = 2 + (k % 5) * 3
        r0 = 5 + (k // 5) * 5
        col, agg = _col(df, kp["column"]), _agg(kp.get("agg", "sum"))
        if agg in fn:
            val = "=" + fn[agg].format(r=rng(col))
        else:  # distintos: valor calculado agora (não há fórmula simples em todas as versões)
            val = int(df[col].nunique()) if agg == "nunique" else float(df[col].agg(agg))
        dash.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c0 + 1)
        dash.merge_cells(start_row=r0 + 1, start_column=c0, end_row=r0 + 2, end_column=c0 + 1)
        lab, v = dash.cell(r0, c0, kp.get("label", col)), dash.cell(r0 + 1, c0, val)
        lab.font, lab.alignment = Font(size=10, bold=True, color="FFFFFF"), Alignment(horizontal="center")
        v.font, v.alignment = Font(size=20, bold=True, color=NAVY), Alignment(horizontal="center", vertical="center")
        v.number_format = _fmt(kp.get("format", "decimal" if agg == "mean" else "int"))
        for rr in range(r0, r0 + 3):
            for cc in (c0, c0 + 1):
                cell = dash.cell(rr, cc)
                cell.fill = PatternFill("solid", fgColor=PALETTE[k % len(PALETTE)] if rr == r0 else LIGHT)
                cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    chart_row = 6 + 5 * ((len(kpis or []) + 4) // 5) if kpis else 5
    ccol = 1
    fx = {"sum": '=SUMIF({x},{k},{y})', "count": '=COUNTIF({x},{k})', "mean": '=IFERROR(AVERAGEIF({x},{k},{y}),0)',
          "max": '=_xlfn.MAXIFS({y},{x},{k})', "min": '=_xlfn.MINIFS({y},{x},{k})'}
    for i, ch in enumerate(charts or []):
        x = _col(df, ch["x"])
        agg = _agg(ch.get("agg", "sum" if ch.get("y") else "count"))
        y = _col(df, ch["y"]) if ch.get("y") else x
        if agg not in fx:
            raise ValueError(f"gráfico: agregação '{agg}' não suportada (use sum, count, mean, max, min)")
        g = df.groupby(x)[y].agg("count" if agg == "count" else agg) if agg != "sum" else \
            pd.to_numeric(df[y], errors="coerce").groupby(df[x]).sum()
        g = g.sort_index() if ch.get("_sort") == "x" or ch.get("type") in ("line", "area") else \
            g.sort_values(ascending=False)
        g = g.head(int(ch.get("top", 12)))
        calc.cell(1, ccol, x).font = Font(bold=True)
        calc.cell(1, ccol + 1, ch.get("title") or f"{agg} de {y}").font = Font(bold=True)
        for r, key in enumerate(g.index, 2):
            kc = calc.cell(r, ccol, key if not hasattr(key, "item") else key.item())
            ref = f"Resumos!${get_column_letter(ccol)}${r}"
            calc.cell(r, ccol + 1, fx[agg].format(x=rng(x), y=rng(y), k=ref)).number_format = \
                _fmt(ch.get("format", "decimal" if agg == "mean" else "int"))
            kc.number_format = "@"
        calc.column_dimensions[get_column_letter(ccol)].width = 22
        spec = {"type": ch.get("type", "bar"), "title": ch.get("title", f"{y} por {x}"),
                "category_col": ccol, "value_cols": [ccol + 1], "width": 16, "height": 8}
        chart = _chart(calc, spec, len(g) + 1)
        chart.legend = None if spec["type"] != "pie" else chart.legend
        dash.add_chart(chart, f"{get_column_letter(2 + (i % 2) * 8)}{chart_row + (i // 2) * 17}")
        ccol += 3
    wb.calculation.fullCalcOnLoad = True  # Excel recalcula tudo ao abrir
    wb.save(config.safe_path(path))
    return {"path": path, "rows": len(df), "kpis": len(kpis or []), "charts": len(charts or []),
            "sheets": wb.sheetnames}


@tool("excel_update", "Edita uma planilha existente sem perder formatação: cells={\"A1\": valor ou "
      "\"=FÓRMULA\"}, append_rows=[[...]] adiciona no fim, new_sheet cria aba com linhas.",
      {"path": {"type": "string"}, "sheet": {"type": "string"}, "cells": {"type": "object"},
       "append_rows": {"type": "array"}, "new_sheet": {"type": "string"}, "rows": {"type": "array"}},
      ("path",))
def excel_update(path: str, sheet: str | None = None, cells: dict | None = None,
                 append_rows: list | None = None, new_sheet: str = "", rows: list | None = None):
    p = config.safe_path(path)
    wb = load_workbook(p, keep_vba=p.suffix.lower() == ".xlsm")
    if new_sheet:
        ws = wb[new_sheet] if new_sheet in wb.sheetnames else wb.create_sheet(new_sheet[:31])
        if rows:
            _fill_sheet(ws, rows)
    ws = wb[sheet] if sheet else wb.active
    for ref, v in (cells or {}).items():
        ws[ref] = v
    for r in append_rows or []:
        ws.append(list(r))
    wb.save(p)
    return {"path": path, "sheet": ws.title, "rows": ws.max_row}


@tool("excel_read", "Lê uma planilha .xlsx/.csv do workspace e retorna as linhas (para analisar, "
      "prefira data_profile/data_query).",
      {"path": {"type": "string"}, "sheet": {"type": "string"}, "max_rows": {"type": "integer"}},
      ("path",))
def excel_read(path: str, sheet: str | None = None, max_rows: int = 500):
    p = config.safe_path(path)
    if p.suffix.lower() == ".csv":
        import csv
        with open(p, newline="", encoding="utf-8-sig") as f:
            return list(csv.reader(f))[:max_rows]
    ws = load_workbook(p, data_only=True)[sheet] if sheet else load_workbook(p, data_only=True).active
    return [list(r) for r in ws.iter_rows(values_only=True, max_row=max_rows)]


@tool("excel_refresh", "Abre no Excel (Windows) e atualiza tudo: consultas Power Query, conexões e "
      "tabelas dinâmicas; opcionalmente roda uma macro; salva. Macros pedem autorização.",
      {"path": {"type": "string"}, "macro": {"type": "string"}}, ("path",),
      risky=lambda a: bool(a.get("macro")))
def excel_refresh(path: str, macro: str = ""):
    if sys.platform != "win32":
        raise RuntimeError("excel_refresh usa o Excel instalado: só no Windows.")
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible, xl.DisplayAlerts = False, False
    try:
        wb = xl.Workbooks.Open(str(config.safe_path(path)))
        wb.RefreshAll()
        xl.CalculateUntilAsyncQueriesDone()
        if macro:
            xl.Application.Run(f"'{wb.Name}'!{macro}")
        wb.Save()
        wb.Close(False)
    finally:
        xl.Quit()
    return {"refreshed": path, "macro": macro or None}
