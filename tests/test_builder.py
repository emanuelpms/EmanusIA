"""Construtor sem código: PDF, XML/NF-e, pastas, e-mail com filtros, expressões, catálogo e teste ao vivo."""
import email.message
import os
import time
import zipfile

import pytest
import yaml
from fastapi.testclient import TestClient

from emanusia import auth, catalog, config
from emanusia.engine import safe_eval, validate
from emanusia.tools import registry
from emanusia.tools.workflow import run_steps

H = {"host": "127.0.0.1"}
ADMIN = auth.LOCAL_ADMIN


def make_pdf(path, pages: list[list[str]]):
    """PDF mínimo com texto de verdade (Helvetica), uma lista de linhas por página."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /Font /Helvetica /BaseFont /Helvetica >>"]
    kids = []
    for lines in pages:
        esc = lambda t: t.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = "BT /F1 11 Tf 50 780 Td 14 TL " + " ".join(f"({esc(l)}) Tj T*" for l in lines) + " ET"
        objs.append(f"<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}\nendstream")
        content = len(objs)
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents {content} 0 R "
                    f"/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(len(objs))
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>"
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode() + b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    p = config.safe_path(path)
    p.write_bytes(out)
    return p


NFE = """<?xml version="1.0" encoding="UTF-8"?>
<nfeProc xmlns="http://www.portalfiscal.inf.br/nfe" versao="4.00"><NFe><infNFe Id="NFe3526{n:0>40}" versao="4.00">
<ide><nNF>{n}</nNF><serie>1</serie><dhEmi>2026-03-0{n}T10:00:00-03:00</dhEmi><natOp>Venda</natOp></ide>
<emit><CNPJ>11222333000181</CNPJ><xNome>Fornecedor Alfa</xNome><enderEmit><UF>SP</UF></enderEmit></emit>
<dest><CNPJ>99888777000166</CNPJ><xNome>Cliente Beta</xNome></dest>
<det nItem="1"><prod><cProd>A1</cProd><xProd>Parafuso</xProd><NCM>73181500</NCM><CFOP>5102</CFOP><uCom>UN</uCom>
<qCom>10.0000</qCom><vUnCom>1.50</vUnCom><vProd>15.00</vProd></prod></det>
<det nItem="2"><prod><cProd>B2</cProd><xProd>Porca</xProd><NCM>73181600</NCM><CFOP>5102</CFOP><uCom>UN</uCom>
<qCom>4.0000</qCom><vUnCom>2.50</vUnCom><vProd>10.00</vProd></prod></det>
<total><ICMSTot><vProd>25.00</vProd><vFrete>0.00</vFrete><vDesc>0.00</vDesc><vICMS>4.50</vICMS><vIPI>0.00</vIPI>
<vPIS>0.16</vPIS><vCOFINS>0.76</vCOFINS><vNF>{total}</vNF></ICMSTot></total></infNFe></NFe></nfeProc>"""


# ---------------------------------------------------------------- PDF
def test_pdf_text_find_fields_merge_split():
    make_pdf("docs/fatura.pdf", [["Fatura 123", "CNPJ: 11.222.333/0001-81", "Vencimento: 10/04/2026", "Valor total: R$ 1.234,56"],
                                 ["Obrigado", "contato@fornecedor.com"]])
    t = registry.call("pdf_text", {"path": "docs/fatura.pdf"})
    assert t["pages"] == 2 and "Valor total" in t["text"]
    assert registry.call("pdf_find", {"path": "docs/fatura.pdf", "find": "cnpj"}) == ["11.222.333/0001-81"]
    assert registry.call("pdf_find", {"path": "docs/fatura.pdf", "find": "valor", "first": True}) == "R$ 1.234,56"
    assert registry.call("pdf_find", {"text": t["text"], "find": "email"}) == ["contato@fornecedor.com"]
    f = registry.call("pdf_fields", {"path": "docs/fatura.pdf", "fields": {"venc": "Vencimento", "total": "Valor total", "x": "Nada"}})
    assert f == {"venc": "10/04/2026", "total": "R$ 1.234,56", "x": ""}
    assert registry.call("pdf_info", {"path": "docs/fatura.pdf"})["pages"] == 2
    make_pdf("docs/outra.pdf", [["Outra"]])
    m = registry.call("pdf_merge", {"files": "docs/*.pdf", "output": "juntos.pdf"})
    assert m["files"] == 2 and registry.call("pdf_info", {"path": "juntos.pdf"})["pages"] == 3
    parts = registry.call("pdf_split", {"path": "juntos.pdf", "folder": "partes"})
    assert parts == ["partes/juntos_p1.pdf", "partes/juntos_p2.pdf", "partes/juntos_p3.pdf"]
    assert "Obrigado" in registry.call("pdf_text", {"path": "partes/juntos_p2.pdf"})["text"]


# ---------------------------------------------------------------- XML
def test_nfe_read_folder_to_table_and_items():
    for n, total in ((1, "25.00"), (2, "1500.10")):
        config.safe_path(f"notas/n{n}.xml").write_text(NFE.format(n=n, total=total), encoding="utf-8")
    config.safe_path("notas/outro.xml").write_text("<pedido><id>1</id></pedido>", encoding="utf-8")
    r = registry.call("nfe_read", {"files": "notas/*.xml", "output": "notas.xlsx"})
    head = r["rows"][0]
    row = dict(zip(head, r["rows"][1]))
    assert r["total_rows"] == 2 and r["notes"] == 2 and r["not_nfe"] == ["outro.xml"]
    assert row["numero"] == "1" and row["emitente"] == "Fornecedor Alfa" and row["valor_total"] == 25.0
    assert row["chave"].startswith("3526") and len(row["chave"]) == 44 and row["emissao"] == "2026-03-01"
    assert (config.workspace() / "notas.xlsx").exists()
    it = registry.call("nfe_read", {"files": ["notas/n1.xml"], "items": True})
    assert it["total_rows"] == 2 and dict(zip(it["rows"][0], it["rows"][2]))["descricao"] == "Porca"


def test_xml_read_items_and_dict():
    config.safe_path("p.xml").write_text("<lote><pedido id='7'><cliente><nome>Ana</nome></cliente><valor>10.5</valor></pedido>"
                                         "<pedido id='8'><cliente><nome>Bia</nome></cliente><valor>3</valor></pedido></lote>")
    r = registry.call("xml_read", {"files": "p.xml", "items": "pedido", "fields": {"id": "@id", "cliente": "cliente/nome", "valor": "valor"}})
    assert r["rows"] == [["id", "cliente", "valor", "arquivo"], ["7", "Ana", 10.5, "p.xml"], ["8", "Bia", 3.0, "p.xml"]]
    d = registry.call("xml_read", {"files": "p.xml"})
    assert d["pedido"][1]["cliente"]["nome"] == "Bia"


def test_xml_blocks_external_entities():
    config.safe_path("x.xml").write_text('<?xml version="1.0"?><!DOCTYPE r [<!ENTITY e SYSTEM "file:///etc/passwd">]><r><a>&e;</a></r>')
    d = registry.call("xml_read", {"files": "x.xml"})
    assert "root:" not in str(d)


# ---------------------------------------------------------------- pastas e arquivos
def test_folder_operations_and_allowed_folders(tmp_path):
    for n in ("a.pdf", "b.pdf", "c.txt"):
        config.safe_path(f"entrada/{n}").write_text(n)
    assert registry.call("file_list", {"subdir": "entrada", "pattern": "*.pdf"}) == ["entrada/a.pdf", "entrada/b.pdf"]
    assert registry.call("file_copy", {"files": "entrada/*.pdf", "destination": "copia"}) == ["copia/a.pdf", "copia/b.pdf"]
    assert registry.call("file_rename", {"path": "copia/a.pdf", "new_name": "nota_1.pdf"}) == "copia/nota_1.pdf"
    with pytest.raises(ValueError):
        registry.call("file_rename", {"path": "copia/b.pdf", "new_name": "../fora.pdf"})
    z = registry.call("file_zip", {"files": "copia/*.pdf", "output": "pacote.zip"})
    assert z["files"] == 2
    assert sorted(registry.call("file_unzip", {"path": "pacote.zip", "folder": "unz"})) == ["unz/b.pdf", "unz/nota_1.pdf"]
    assert registry.call("file_move", {"files": "entrada/*.pdf", "destination": "feitos/"}) == ["feitos/a.pdf", "feitos/b.pdf"]
    assert registry.call("file_exists", {"path": "feitos/a.pdf"}) and not registry.call("file_exists", {"path": "entrada/a.pdf"})
    registry.call("file_write", {"path": "log.txt", "content": "l1", "append": True})
    registry.call("file_write", {"path": "log.txt", "content": "l2", "append": True})
    assert registry.call("file_read", {"path": "log.txt"}) == "l1\nl2\n"
    registry.call("json_write", {"path": "d.json", "data": [{"a": 1}]})
    assert registry.call("json_read", {"path": "d.json"}) == [{"a": 1}]
    assert registry.get("file_delete").is_risky({})
    # pasta fora do workspace só depois de liberada pelo admin
    outside = tmp_path / "Notas Fiscais"
    outside.mkdir()
    (outside / "n.txt").write_text("ok")
    with pytest.raises(PermissionError, match="Pasta não liberada"):
        registry.call("file_list", {"subdir": str(outside)})
    config.save_settings(allowed_folders=[str(outside)])
    assert registry.call("file_list", {"subdir": str(outside)}) == [str((outside / "n.txt").resolve())]
    with pytest.raises(PermissionError):
        registry.call("file_read", {"path": str(tmp_path / "data" / "users.json")})


def test_zip_slip_blocked():
    p = config.safe_path("mal.zip")
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("../../fora.txt", "x")
    with pytest.raises(PermissionError):
        registry.call("file_unzip", {"path": "mal.zip", "folder": "dest"})


# ---------------------------------------------------------------- e-mail
def test_email_read_filters_imap(monkeypatch):
    from emanusia.tools import mail
    msgs = []
    for subj, sender, att in (("Nota fiscal 1", "fornecedor@alfa.com", ("n1.xml", b"<x/>")),
                              ("Reunião", "chefe@empresa.com", None),
                              ("Nota fiscal 2", "fornecedor@alfa.com", ("boleto.pdf", b"%PDF"))):
        m = email.message.EmailMessage()
        m["Subject"], m["From"], m["Date"] = subj, sender, "Mon, 02 Mar 2026 10:00:00 -0300"
        m.set_content("corpo")
        if att:
            m.add_attachment(att[1], maintype="application", subtype="octet-stream", filename=att[0])
        msgs.append(m.as_bytes())

    class FakeIMAP:
        def __init__(self, host): pass
        def __enter__(self): return self
        def __exit__(self, *a): pass
        def login(self, u, p): pass
        def select(self, f, readonly=True): pass
        def search(self, charset, *crit):
            FakeIMAP.crit = crit
            return "OK", [b" ".join(str(i + 1).encode() for i in range(len(msgs)))]
        def fetch(self, i, what): return "OK", [(b"", msgs[int(i) - 1])]
    monkeypatch.setattr(mail.imaplib, "IMAP4_SSL", FakeIMAP)
    monkeypatch.setenv("EMAIL_IMAP_PASS", "x")
    config.save_settings(email_provider="smtp", imap_host="imap.x.com", imap_user="bot@x.com")
    r = registry.call("email_read", {"subject_contains": "nota", "attachment_types": "xml", "attachments_folder": "xmls",
                                     "since_days": 2})
    assert [e["subject"] for e in r] == ["Nota fiscal 1"] and r[0]["attachments"] == ["xmls/n1.xml"]
    assert "SINCE" in FakeIMAP.crit and (config.workspace() / "xmls/n1.xml").read_bytes() == b"<x/>"
    r = registry.call("email_read", {"from_contains": "alfa", "unread_only": False})
    assert len(r) == 2 and all(not e["attachments"] for e in r)


def test_email_send_attachment_pattern_and_missing(monkeypatch):
    from emanusia.tools import mail
    sent = {}
    monkeypatch.setattr(mail, "outlook_send", lambda to, s, b, att, cc, html: sent.update(att=att))
    config.safe_path("saida/a.pdf").write_text("x")
    config.safe_path("saida/b.pdf").write_text("x")
    registry.call("email_send", {"to": "a@x.com", "subject": "s", "body": "b", "attachments": "saida/*.pdf"})
    assert [p.replace("\\", "/").split("/")[-1] for p in sent["att"]] == ["a.pdf", "b.pdf"]
    with pytest.raises(FileNotFoundError):
        registry.call("email_send", {"to": "a@x.com", "subject": "s", "body": "b", "attachments": ["nao.pdf"]})


# ---------------------------------------------------------------- expressões e Excel
def test_expression_helpers():
    ctx = {"t": [["nome", "valor"], ["a", "R$ 10,50"], ["b", "1.000,00"]], "d": "2026-01-31"}
    assert safe_eval("to_number('R$ 1.234,56')", ctx) == 1234.56
    assert safe_eval("column(t, 'valor')", ctx) == ["R$ 10,50", "1.000,00"]
    assert safe_eval("records(t)[1]['nome']", ctx) == "b"
    assert safe_eval("date_add(d, 0, 1)", ctx) == "2026-02-28" and safe_eval("format_date(d)", ctx) == "31/01/2026"
    assert safe_eval("money(1234.5)", ctx) == "R$ 1.234,50" and safe_eval("days_between('01/01/2026', d)", ctx) == 30
    assert safe_eval("regex_find('Pedido 4471 ok', '\\\\d+')", ctx) == "4471"
    assert safe_eval("join(split('a;b', ';'), ', ')", ctx) == "a, b" and safe_eval("zfill(7, 3)", ctx) == "007"
    assert len(safe_eval("today()", ctx)) == 10


def test_excel_save_formats_totals_chart_and_append():
    from openpyxl import load_workbook
    data = {"rows": [["cliente", "valor"], ["Ana", 10.5], ["Bia", 20]], "total_rows": 2}
    r = registry.call("excel_save", {"path": "r.xlsx", "data": data, "formats": {"valor": "money"}, "totals": ["valor"],
                                     "chart_type": "bar", "chart_x": "cliente", "chart_y": ["valor"]})
    assert r["rows"] == 2
    ws = load_workbook(config.workspace() / "r.xlsx")["Dados"]
    assert ws["B2"].number_format.startswith("R$") and ws["B4"].value == "=SUM(B2:B3)" and ws._charts
    registry.call("excel_save", {"path": "a.xlsx", "data": [{"x": 1, "y": 2}]})
    r = registry.call("excel_save", {"path": "a.xlsx", "data": [{"y": 5, "x": 4}], "append": True})
    ws = load_workbook(config.workspace() / "a.xlsx")["Dados"]
    assert r["added"] == 1 and [c.value for c in ws[3]] == [4, 5]


# ---------------------------------------------------------------- catálogo e modelos
def test_catalog_actions_match_tools_and_templates_validate():
    c = catalog.catalog()
    ids = [a["id"] for a in c["actions"]]
    assert len(ids) == len(set(ids)) and {"pdf", "xml", "email", "excel", "flow"} <= {x["id"] for x in c["categories"]}
    for a in catalog.ACTIONS:
        if a.get("tool"):
            t = registry.get(a["tool"])
            for f in a["fields"]:
                for k in [f.get("key") or "args." + f["name"]] if not isinstance(f.get("key"), list) else f["key"]:
                    if k.startswith("args."):
                        assert k[5:] in t.parameters, (a["id"], k)
    for t in catalog.TEMPLATES:
        rep = validate(t["code"])
        assert rep["ok"], (t["id"], rep["errors"])
    import re

    def sub(v, vals):
        if isinstance(v, str):
            return re.sub(r"\$(\w+)", lambda m: str(vals[m.group(1)]), v)
        if isinstance(v, list):
            return [sub(x, vals) for x in v]
        if isinstance(v, dict):
            return {sub(k, vals): sub(x, vals) for k, x in v.items()}
        return v
    for a in catalog.ACTIONS:  # atalhos de vários passos também viram código válido
        if a.get("template"):
            vals = {f["name"]: f.get("default") or "x" for f in a["fields"]}
            rep = validate({"steps": sub(a["template"], vals)})
            assert rep["ok"], (a["id"], rep["errors"])


@pytest.fixture
def client():
    from emanusia import api
    with TestClient(api.app, headers=H) as c:
        yield c


def test_builder_check_maps_problem_paths(client):
    code = {"steps": [{"tool": "pdf_text", "args": {}},
                      {"if": "x > 1", "then": [{"log": "ok"}], "else": [{"tool": "web_click"}]},
                      {"try": [{"tool": "nfe_read"}], "catch": [{"tool": "nao_existe"}]}]}
    r = client.post("/api/builder/check", json={"code": code}).json()
    paths = {p["path"]: p["text"] for p in r["problems"]}
    assert set(paths) == {"1", "2.else.1", "3.1", "3.catch.1"}
    assert "falta informar" in paths["1"] and "não sabe fazer" in paths["3.catch.1"]
    assert [s["path"] for s in r["steps"]][:4] == ["1", "2", "2.1", "2.else"]
    assert "pdf_text" in r["yaml"]
    cat = client.get("/api/catalog").json()
    assert cat["templates"] and cat["step_options"]


def test_builder_bot_rehearsal_reports_each_step(client):
    config.safe_path("entrada/a.txt").write_text("um")
    config.safe_path("entrada/b.txt").write_text("dois")
    code = {"params": {"pasta": "entrada"}, "steps": [
        {"tool": "file_list", "args": {"subdir": "{{pasta}}", "pattern": "*.txt"}, "save_as": "arquivos"},
        {"set": {"linhas": []}},
        {"for_each": "arquivos", "as": "arquivo", "steps": [
            {"tool": "file_read", "args": {"path": "{{arquivo}}"}, "save_as": "texto"},
            {"set": {"linhas": "=linhas + [{'arquivo': arquivo, 'texto': texto}]"}}]},
        {"tool": "excel_save", "args": {"path": "saida.xlsx", "data": "{{linhas}}"}, "save_as": "planilha"}]}
    b = client.post("/api/bots", json={"name": "Leitor de pasta", "code": code, "origin": "construtor"}).json()
    assert b["status"] == "validated" and b["origin"] == "construtor"
    jid = client.post(f"/api/bots/{b['id']}/rehearse", json={"params": {}}).json()["id"]
    for _ in range(100):
        j = client.get(f"/api/jobs/{jid}").json()
        if j["status"] in ("done", "error"):
            break
        time.sleep(0.1)
    assert j["status"] == "done" and j["result"]["ok"], j
    done = [e for e in j["events"] if e["kind"] == "step_done"]
    assert [e["path"] for e in done] == ["1", "3.1", "3.1", "4"]
    assert done[0]["preview"]["text"].startswith("Lista com 2") and done[-1]["secs"] >= 0
    assert (config.workspace() / "saida.xlsx").exists()
    assert client.get("/builder").status_code == 200


def test_settings_allowed_folders_admin_only(client, tmp_path):
    r = client.post("/api/settings", json={"allowed_folders": [str(tmp_path), "  "]}).json()
    assert r["allowed_folders"] == [str(tmp_path)]


# ---------------------------------------------------------------- tela do construtor (navegador de verdade)
def test_builder_ui_build_save_and_test(headless):
    import socket
    import threading

    import uvicorn
    from playwright.sync_api import sync_playwright
    for n, total in ((1, "25.00"), (2, "99.90")):
        config.safe_path(f"notas/n{n}.xml").write_text(NFE.format(n=n, total=total), encoding="utf-8")
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config("emanusia.api:app", host="127.0.0.1", port=port, log_level="warning"))
    threading.Thread(target=server.run, daemon=True).start()
    while not server.started:
        time.sleep(0.05)
    errs = []
    try:
        with sync_playwright() as pw:
            kw = {"executable_path": os.environ["EMANUSIA_CHROMIUM"]} if os.environ.get("EMANUSIA_CHROMIUM") else {}
            b = pw.chromium.launch(**kw)
            pg = b.new_page(viewport={"width": 1400, "height": 900})
            pg.on("pageerror", lambda e: errs.append(str(e)))
            pg.on("dialog", lambda d: d.accept())
            pg.goto(f"http://127.0.0.1:{port}/builder")
            pg.wait_for_selector(".tpl")
            pg.fill("#bn", "Notas do dia")
            pg.fill("#q", "nf-e")
            pg.click("button.act:has-text('Ler notas fiscais')")
            pg.locator("#props input").first.fill("notas/*.xml")
            pg.fill("#q", "planilha")
            pg.click("button.act:has-text('Montar planilha com os dados')")
            pg.locator("#props input").first.fill("relatorio.xlsx")
            pg.locator("#props .vb").nth(1).click()
            pg.click("#pop button:has-text('notas')")
            pg.click("#testb")
            pg.wait_for_selector(".banner.ok, .banner.bad", timeout=30000)
            banner = pg.inner_text(".banner")
            flow = pg.inner_text("#flow")
            b.close()
    finally:
        server.should_exit = True
    assert not errs and "deu tudo certo" in banner, (errs, banner)
    assert "Tabela com 2 linha(s)" in flow and "Salvo em relatorio.xlsx" in flow
    from emanusia import bots
    bt = bots.get_bot("notas-do-dia")
    assert bt.origin == "construtor" and bt.rehearsal["ok"]
    assert bt.code["steps"][1] == {"tool": "excel_save", "args": {"path": "relatorio.xlsx", "data": "{{notas}}", "sheet": "Dados",
                                                                  "as_table": True}}


# ---------------------------------------------------------------- analista + processo completo
def _base_vendas(n_extra_dup=2):
    import random
    random.seed(3)
    rows = [["data", "regiao", "vendedor", "valor"]]
    for m, peso in ((8, 1.0), (9, 1.1), (10, 0.7)):  # outubro cai
        for i in range(40):
            reg = random.choice(["Sul", "Sudeste", "Sudeste", "Sudeste", "Norte", "Nordeste", "Centro-Oeste"])
            if m == 10 and reg == "Sudeste":
                peso_r = 0.5
            else:
                peso_r = 1
            v = round(random.uniform(200, 900) * peso * peso_r, 2)
            rows.append([f"{random.randint(1, 28):02d}/{m:02d}/2026", reg, random.choice(["Ana", "Rui", "Lia"]),
                         f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")])
    rows.append(["15/10/2026", "Sul", "Ana", "R$ 48.000,00"])  # fora do padrão
    rows += [rows[1]] * n_extra_dup
    return rows


def test_data_insights_writes_findings_and_suggestions():
    registry.call("excel_write", {"path": "vendas.xlsx", "sheets": {"Base": _base_vendas()}})
    dados = registry.call("data_clean", {"source": "vendas.xlsx", "dtypes": {"valor": "number", "data": "date"}})
    assert dados["total_rows"] == 121 and len(dados["rows"]) == 21  # na tela, amostra...
    r = registry.call("data_insights", {"source": dados, "value": "valor", "date_column": "data",
                                        "dimensions": ["regiao"], "goal": 100000, "title": "Vendas"})
    txt = r["text"]
    assert r["summary"]["registros"] == 121  # ...mas a análise usa a base inteira
    assert "out/2026" in txt and "em relação a set/2026" in txt
    assert "Sem os valores fora do padrão, a variação de out/2026 seria -" in txt and "Investigar a queda" in txt
    assert "Maior queda por regiao no último mês: Sudeste" in txt
    assert "muito acima do normal" in txt and "R$ 48.000,00" in txt
    assert "Meta de R$ 100.000,00" in txt
    assert r["html"].startswith("<h3") and "<li>" in r["html"] and "Sugestões de ação" in r["html"]
    vazio = registry.call("data_insights", {"source": [["valor"]], "value": "valor"})
    assert "vazia" in vazio["text"]


def test_full_process_template_web_excel_dashboard_desktop_emails(headless, monkeypatch):
    import functools
    import http.server
    import threading

    from emanusia import vault
    from emanusia.tools import mail
    ws = config.workspace()
    registry.call("excel_write", {"path": "site/base.xlsx", "sheets": {"Base": _base_vendas()}})
    (ws / "site/index.html").write_text("""<meta charset=utf-8><form onsubmit="event.preventDefault();document.getElementById('app').hidden=false;this.hidden=true">
      <input type=text name=u><input type=password name=p><button>Entrar</button></form>
      <div id=app hidden><button onclick="document.getElementById('dl').hidden=false">Relatórios</button>
      <a id=dl hidden href="base.xlsx" download>Exportar Excel</a></div>""", encoding="utf-8")
    handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(ws / "site"))
    handler.log_message = lambda *a: None
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    vault.set_credential("local", "sistema-web", "robo", "segredo", url="")
    typed, sent = [], []
    for name in ("desktop_start_app", "desktop_type", "desktop_click"):  # "outro programa do PC" (Windows)
        monkeypatch.setattr(registry.get(name), "fn", lambda _n=name, **a: typed.append((_n, a)) or "ok")
    monkeypatch.setattr(mail, "outlook_send", lambda to, subj, body, att, cc, html: sent.append((to, subj, body, att, html)))
    t = next(x for x in catalog.TEMPLATES if x["id"] == "processo_completo")
    try:
        run_steps({**t["code"], "name": ""}, {"site": f"http://127.0.0.1:{srv.server_address[1]}/index.html"},
                  confirm=lambda n, a: True)
    finally:
        srv.shutdown()
        registry.call("web_close")
    for f in ("downloads/base.xlsx", "saida/dados-tratados.xlsx", "saida/dashboard.xlsx", "saida/painel.html",
              "saida/relatorio-lancamentos.xlsx"):
        assert (ws / f).exists(), f
    from openpyxl import load_workbook
    assert load_workbook(ws / "saida/dados-tratados.xlsx").active.max_row == 122  # base inteira tratada
    regioes = [a["text"] for n, a in typed if n == "desktop_type" and a["control"] == "Região"]
    assert sorted(regioes) == ["Centro-Oeste", "Nordeste", "Norte", "Sudeste", "Sul"]
    assert [s[0] for s in sent] == ["ti@empresa.com", "gestores@empresa.com"]
    assert "5 lançamento(s)" in sent[0][2]
    to, subj, body, att, html_on = sent[1]
    assert html_on and "O que foi identificado" in body and "Sugestões de ação" in body and len(att) == 5


def test_full_process_failure_warns_responsible(monkeypatch):
    from emanusia.tools import mail
    sent = []
    monkeypatch.setattr(mail, "outlook_send", lambda to, subj, body, att, cc, html: sent.append((to, subj, body)))
    t = next(x for x in catalog.TEMPLATES if x["id"] == "processo_completo")
    code = {**t["code"], "name": ""}
    code["steps"] = [{**code["steps"][0], "try": [{"tool": "data_clean", "args": {"source": "nao-existe.xlsx"}}]}]
    with pytest.raises(Exception):
        run_steps(code, confirm=lambda n, a: True)
    assert sent and sent[0][1].startswith("⚠️") and "nao-existe" in sent[0][2]
