"""Tudo dentro do sistema: fórmulas calculadas sem o Excel, visualizador de planilhas/PDF/painéis."""
import datetime as dt

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from emanusia import config, xlformula
from emanusia.tools import registry

H = {"host": "127.0.0.1"}


def _wb(path, sheets: dict):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    p = config.safe_path(path)
    wb.save(p)
    return p


def test_formula_engine_common_functions():
    p = _wb("f.xlsx", {
        "Dados": [["regiao", "valor", "qtd"], ["Sul", 10, 1], ["Norte", 20.5, 2], ["Sul", 5, 3], ["Leste", None, 4]],
        "Calc": [["=SUM(Dados!B:B)", "=SUMIF(Dados!A:A,\"Sul\",Dados!B:B)", "=COUNTIF(Dados!A:A,\"S*\")",
                  "=AVERAGEIF(Dados!A:A,\"Sul\",Dados!B:B)", "=COUNTA(Dados!A:A)-1", "=MAX(Dados!B2:B5)"],
                 ["=SUMIFS(Dados!B:B,Dados!A:A,\"Sul\",Dados!C:C,\">1\")", "=COUNTIFS(Dados!A:A,\"<>Sul\")",
                  "=IFERROR(1/0,\"sem\")", "=1/0", "=VLOOKUP(\"Norte\",Dados!A2:C5,3,FALSE)", "=ROUND(2.345,2)"],
                 ["=IF(A1>30,\"alto\",\"baixo\")", "=\"Total: \"&A1", "=_xlfn.MAXIFS(Dados!B:B,Dados!A:A,\"Sul\")",
                  "=E3+1", "=D3+1", "=SOMA(1)"],
                 ["=A1*10%", "=COUNT(Dados!B:B)", "=AND(A1>1,B1<20)", "=Dados!B3*2", "=VLOOKUP(\"X\",Dados!A2:C5,2,FALSE)",
                  "=IFERROR(E4,0)"]]})
    v = xlformula.evaluate(p)["Calc"]
    get = lambda ref: v[(int(ref[1:]), ord(ref[0]) - 64)]  # noqa: E731
    assert get("A1") == 35.5 and get("B1") == 15 and get("C1") == 2 and get("D1") == 7.5 and get("E1") == 4
    assert get("F1") == 20.5 and get("A2") == 5 and get("B2") == 3 and get("C2") == "sem" and get("D2") == xlformula.DIV0
    assert get("E2") == 2 and get("F2") == 2.35 and get("A3") == "alto" and get("B3") == "Total: 35.5"
    assert get("C3") == 10 and get("D3") == xlformula.REF and get("E3") == xlformula.REF  # circular
    assert get("F3") == xlformula.NAME  # função desconhecida
    assert get("A4") == pytest.approx(3.55) and get("B4") == 3 and get("C4") is True and get("D4") == 41
    assert get("E4") == xlformula.NA and get("F4") == 0


def test_generated_excel_files_carry_calculated_values():
    data = [["data", "regiao", "valor"], ["05/03/2026", "Sul", 100.0], ["06/03/2026", "Norte", 50.5],
            ["10/04/2026", "Sul", 25.0]]
    registry.call("excel_dashboard", {"path": "d.xlsx", "source": data, "title": "Vendas",
                                      "kpis": [{"label": "Total", "column": "valor", "agg": "sum", "format": "money"}],
                                      "charts": [{"type": "bar", "x": "regiao", "y": "valor", "agg": "sum"}]})
    wb = load_workbook(config.workspace() / "d.xlsx", data_only=True)  # como qualquer leitor sem Excel veria
    assert wb["Dashboard"]["B6"].value == 175.5
    assert [wb["Resumos"].cell(r, 2).value for r in (2, 3)] == [125.0, 50.5]
    assert load_workbook(config.workspace() / "d.xlsx")["Dashboard"]["B6"].value.startswith("=SUM(")  # fórmula continua
    registry.call("excel_save", {"path": "r.xlsx", "data": data, "totals": ["valor"]})
    df = pd.read_excel(config.workspace() / "r.xlsx")
    assert df.iloc[-1]["valor"] == 175.5  # linha de TOTAL legível pelos próximos passos


def test_excel_save_dates_and_summary_chart():
    d = registry.call("data_clean", {"source": [["data", "regiao", "valor"], ["05/03/2026", "Sul", "R$ 10,00"],
                                                ["06/03/2026", "Sul", "R$ 5,00"], ["07/03/2026", "Norte", "R$ 1,00"]],
                                     "dtypes": {"data": "date", "valor": "number"}})
    registry.call("excel_save", {"path": "g.xlsx", "data": d, "chart_type": "bar", "chart_x": "regiao", "chart_y": ["valor"]})
    wb = load_workbook(config.workspace() / "g.xlsx", data_only=True)
    c = wb["Dados"]["A2"]
    assert isinstance(c.value, dt.datetime) and c.number_format == "dd/mm/yyyy"
    assert [[x.value for x in r] for r in wb["Resumo"].iter_rows()] == [["regiao", "valor"], ["Sul", 15], ["Norte", 1]]


def test_excel_recalc_tool_reports_problems_and_values_copy():
    _wb("m.xlsx", {"P": [[1, 2, "=A1+B1", "=FUNCAOX(1)"]]})
    r = registry.call("excel_recalc", {"path": "m.xlsx", "values_only": True})
    assert r["formulas"] == 2 and r["problems"] == ["P!D1: #NAME?"] and r["output"] == "m-valores.xlsx"
    assert load_workbook(config.workspace() / "m-valores.xlsx")["P"]["C1"].value == 3  # sem fórmula


@pytest.fixture
def client():
    from emanusia import api
    with TestClient(api.app, headers=H) as c:
        yield c


def test_viewer_sheet_with_formats_merges_and_charts(client):
    registry.call("excel_dashboard", {"path": "saida/d.xlsx", "title": "Painel", "source": [
        ["regiao", "valor"], ["Sul", 1234.5], ["Norte", 10]],
        "kpis": [{"label": "Total", "column": "valor", "agg": "sum", "format": "money"}],
        "charts": [{"type": "pie", "x": "regiao", "y": "valor", "title": "Por região"}]})
    v = client.get("/api/view", params={"path": "saida/d.xlsx"}).json()
    assert v["kind"] == "sheet" and [s["name"] for s in v["sheets"]] == ["Dashboard", "Dados", "Resumos"]
    dash = v["sheets"][0]
    assert dash["rows"][5][1]["v"] == "R$ 1.244,50" and dash["rows"][1][1]["b"] == 1
    assert [2, 2, 2, 16] in dash["merges"]
    ch = dash["charts"][0]
    assert ch["type"] == "pie" and ch["title"] == "Por região" and ch["series"][0]["x"] == ["Sul", "Norte"]
    assert ch["series"][0]["y"] == [1234.5, 10] and ch["anchor"]["from"]["col"] >= 1
    assert client.get("/api/view", params={"path": "../fora.xlsx"}).status_code == 403
    assert client.get("/api/view", params={"path": "nao.xlsx"}).status_code == 404
    assert client.get("/view").status_code == 200 and client.get("/lib/plotly.min.js").status_code == 200


def test_viewer_pdf_pages_csv_text_and_downloads(client):
    from test_builder import make_pdf
    make_pdf("doc.pdf", [["Pagina um"], ["Pagina dois"]])
    v = client.get("/api/view", params={"path": "doc.pdf"}).json()
    assert v["kind"] == "pdf" and v["pages"] == 2
    png = client.get("/api/view/pdf", params={"path": "doc.pdf", "page": 2})
    assert png.status_code == 200 and png.content[:4] == b"\x89PNG"
    assert client.get("/api/view/pdf", params={"path": "doc.pdf", "page": 9}).status_code == 404
    config.safe_path("t.csv").write_text("a;b\n1;2\n", encoding="utf-8")
    c = client.get("/api/view", params={"path": "t.csv"}).json()
    assert [[x["v"] for x in r] for r in c["sheets"][0]["rows"]] == [["a", "b"], ["1", "2"]]
    config.safe_path("n.json").write_text('{"a": 1}', encoding="utf-8")
    assert client.get("/api/view", params={"path": "n.json"}).json()["text"] == '{"a": 1}'
    config.safe_path("p.html").write_text("<script>1</script>")
    r = client.get("/files/p.html")
    assert "sandbox" in r.headers["content-security-policy"]
    r = client.get("/files/doc.pdf", params={"download": 1})
    assert "attachment" in r.headers["content-disposition"] and "content-security-policy" not in r.headers


def test_builder_marks_actions_that_need_external_programs(client):
    cat = {a["id"]: a for a in client.get("/api/catalog").json()["actions"]}
    assert cat["excel_refresh"]["needs"] == "Excel instalado" and "SAP" in cat["sap_login"]["needs"]
    assert "needs" not in cat["excel_dashboard"] and "needs" not in cat["excel_recalc"] and "needs" not in cat["pdf_text"]


def test_append_keeps_chart_and_summary_up_to_date(client):
    args = {"path": "c.xlsx", "chart_type": "bar", "chart_x": "a", "chart_y": ["b"], "totals": ["b"]}
    registry.call("excel_save", {**args, "data": [["a", "b"], ["x", 1], ["x", 2]]})
    r = registry.call("excel_save", {**args, "data": [["b", "a"], [3, "y"]], "append": True})
    assert r["added"] == 1 and r["rows"] == 3
    wb = load_workbook(config.workspace() / "c.xlsx", data_only=True)
    assert [[c.value for c in row] for row in wb["Dados"].iter_rows()] == [["a", "b"], ["x", 1], ["x", 2], ["y", 3], ["TOTAL", 6]]
    assert [[c.value for c in row] for row in wb["Resumo"].iter_rows()] == [["a", "b"], ["x", 3], ["y", 3]]
    ch = client.get("/api/view", params={"path": "c.xlsx"}).json()["sheets"][1]["charts"][0]
    assert ch["series"][0]["x"] == ["x", "y"] and ch["series"][0]["y"] == [3, 3]
    registry.call("excel_update", {"path": "c.xlsx", "cells": {"D1": "ok"}})  # editar mantém o gráfico
    assert client.get("/api/view", params={"path": "c.xlsx"}).json()["sheets"][1]["charts"]
