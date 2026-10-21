import json

import pandas as pd
import pytest
import yaml
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from emanusia import auth, bots, config, vault
from emanusia.engine import Runner, validate
from emanusia.tools import registry

H = {"host": "127.0.0.1"}
call = registry.call


# ---------------------------------------------------------------- cofre de acessos
def test_vault_encrypts_and_is_per_user():
    vault.set_credential("ana", "portal", "ana@x.com", "Senha-Ana-1", kind="web", url="https://p")
    vault.set_credential("bob", "portal", "bob@x.com", "Senha-Bob-2")
    raw = (config.data_dir() / "vault.json").read_text()
    assert "Senha-Ana-1" not in raw and "Senha-Bob-2" not in raw  # criptografado em disco
    assert vault.resolve("ana", "portal")["password"] == "Senha-Ana-1"
    assert vault.resolve("bob", "portal")["username"] == "bob@x.com"
    with pytest.raises(KeyError, match="não mapeado"):
        vault.resolve("carla", "portal")
    assert "password" not in json.dumps(vault.list_credentials("ana")).replace("has_password", "")
    # editar sem senha mantém a atual
    vault.set_credential("ana", "portal", "novo@x.com", None)
    assert vault.resolve("ana", "portal")["password"] == "Senha-Ana-1"
    with pytest.raises(ValueError):
        vault.set_credential("ana", "novo", "u", None)  # novo acesso sem senha


def test_vault_shared_fallback_and_own_override():
    vault.set_credential(vault.SHARED, "erp", "servico", "s3rv")
    assert vault.resolve("qualquer", "erp")["username"] == "servico"
    vault.set_credential("ana", "erp", "ana", "propria")
    assert vault.resolve("ana", "erp")["password"] == "propria"


def test_web_login_uses_credentials_of_whoever_runs(headless):
    page = config.workspace() / "login.html"
    page.write_text("""<form onsubmit="event.preventDefault();document.body.innerHTML='<h1 id=ok>Olá '+u.value+
      ' '+(p.value==='S3nha!'?'autorizado':'negado')+'</h1>'"><input id=u type=email><input id=p type=password>
      <button>Entrar</button></form>""")
    vault.set_credential("ana", "portal", "ana@x.com", "S3nha!", url=page.as_uri(),
                         extra={"success_selector": "#ok"})
    vault.set_credential("bob", "portal", "bob@x.com", "errada", url=page.as_uri())
    code = {"steps": [{"tool": "web_login", "args": {"credential": "portal"}},
                      {"tool": "web_text", "args": {"selector": "#ok"}, "save_as": "msg"}]}
    try:
        out_ana = Runner(auth.User("ana", "operator")).run(code)["outputs"]["msg"]
        out_bob = Runner(auth.User("bob", "operator")).run(code)["outputs"]["msg"]
    finally:
        call("web_close", {})
    assert out_ana == "Olá ana@x.com autorizado" and out_bob == "Olá bob@x.com negado"
    audit = (config.data_dir() / "audit.jsonl").read_text()
    assert "S3nha!" not in audit and "errada" not in audit


def test_validation_warns_unmapped_credential():
    rep = validate({"steps": [{"tool": "web_login", "args": {"credential": "portal-x"}}]}, auth.User("ana"))
    assert rep["ok"] and any("não está mapeado" in w for w in rep["warnings"])


def test_api_credentials_never_return_password():
    from emanusia import api
    with TestClient(api.app, headers=H) as c:
        r = c.put("/api/credentials/portal", json={"username": "eu", "password": "Segredo123", "url": "https://x"})
        assert r.status_code == 200 and "Segredo123" not in r.text
        assert "Segredo123" not in c.get("/api/credentials").text
        assert c.put("/api/credentials/novo", json={"username": "eu"}).status_code == 422
        assert c.put("/api/credentials/Nome Inválido", json={"password": "x"}).status_code == 422
        assert c.delete("/api/credentials/portal").status_code == 200


# ---------------------------------------------------------------- dados
VENDAS = [["Data", "Região", "Vendedor", "Produto", "Valor", "Qtd"],
          ["05/01/2026", "Sul", "Ana", "A", 1000.0, 2], ["10/01/2026", "Sul", "Rui", "B", 500.0, 1],
          ["03/02/2026", "Norte", "Ana", "A", 1500.0, 3], ["15/02/2026", "Norte", "Lia", "C", 250.5, 1],
          ["20/03/2026", "Sul", "Rui", "A", 2000.0, 4], ["20/03/2026", "Sul", "Rui", "A", 2000.0, 4]]


def _write_vendas(path="vendas.xlsx"):
    call("excel_write", {"path": path, "sheets": {"Vendas": VENDAS}})
    return path


def test_profile_understands_sheet():
    p = call("data_profile", {"source": _write_vendas()})
    assert p["rows"] == 6 and p["duplicates"] == 1 and p["sheets"] == ["Vendas"]
    valor = next(c for c in p["columns"] if c["column"] == "Valor")
    assert valor["sum"] == 7250.5 and valor["max"] == 2000


def test_query_filter_group_sort_and_save():
    src = _write_vendas()
    r = call("data_query", {"source": src, "where": [{"column": "valor", "op": ">=", "value": 500}],
                            "group_by": ["Região"], "agg": {"Valor": "sum", "Qtd": "sum"},
                            "sort_by": "Valor", "output": "por_regiao.xlsx"})
    assert r["rows"] == [["Região", "Valor", "Qtd"], ["Sul", 5500.0, 11], ["Norte", 1500.0, 3]]
    assert (config.workspace() / "por_regiao.xlsx").exists()
    r = call("data_query", {"source": VENDAS, "where": [{"column": "Data", "op": ">=", "value": "01/02/2026"},
                                                        {"column": "Vendedor", "op": "in", "value": ["Ana", "Lia"]}]})
    assert r["total_rows"] == 2
    with pytest.raises(KeyError, match="Colunas"):
        call("data_query", {"source": VENDAS, "group_by": ["Cidade"]})


def test_pivot_with_totals():
    r = call("data_pivot", {"source": VENDAS, "index": ["Vendedor"], "columns": "Região", "values": "Valor"})
    head, *body = r["rows"]
    assert head == ["Vendedor", "Norte", "Sul", "Total"]
    assert body[-1] == ["Total", 1750.5, 5500.0, 7250.5]


def test_clean_brazilian_numbers_dates_duplicates():
    dirty = [[" Nome ", "Valor", "Data"], [" Ana ", "R$ 1.234,56", "05/01/2026"], [" Ana ", "R$ 1.234,56", "05/01/2026"],
             ["", "", ""], ["Rui", "10,5", "31/12/2025"]]
    r = call("data_clean", {"source": dirty, "dtypes": {"valor": "number", "data": "date"}, "output": "limpo.csv"})
    assert r["removed_rows"] == 2 and r["rows"][1][:2] == ["Ana", 1234.56] and r["rows"][2][1] == 10.5
    assert r["rows"][1][2] == "2026-01-05"


def test_merge_as_procv_and_reconciliation():
    extrato = [["Doc", "Pago"], ["1", 100], ["2", 50], ["9", 70]]
    contas = [["Documento", "Fornecedor"], ["1", "ACME"], ["2", "Beta"], ["3", "Gama"]]
    r = call("data_merge", {"left": contas, "right": extrato, "left_on": "Documento", "right_on": "Doc",
                            "how": "outer", "indicator": True})
    assert r["resumo"] == {"ambos": 2, "só esquerda": 1, "só direita": 1}


def test_combine_many_spreadsheets():
    for f, v in (("filiais/sp.xlsx", 10), ("filiais/rj.xlsx", 20)):
        call("excel_write", {"path": f, "sheets": {"A": [["Item", "Valor"], ["x", v]]}})
    r = call("data_combine", {"pattern": "filiais/*.xlsx", "output": "consolidado.xlsx"})
    assert r["files"] == 2 and r["total_rows"] == 2
    assert sorted(x[2] for x in r["rows"][1:]) == ["filiais/rj.xlsx", "filiais/sp.xlsx"]


def test_compute_columns_safely():
    r = call("data_compute", {"source": VENDAS, "new_columns": {"ticket": "Valor / Qtd",
                                                                "faixa": "'ALTA' if Valor >= 1500 else 'normal'"}})
    assert r["rows"][1][-2:] == [500.0, "normal"] and r["rows"][3][-1] == "ALTA"
    with pytest.raises(Exception):
        call("data_compute", {"source": VENDAS, "new_columns": {"x": "__import__('os')"}})


def test_extract_tables_from_html():
    (config.workspace() / "rel.html").write_text("<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>")
    r = call("data_extract", {"path": "rel.html", "index": 0})
    assert r["rows"] == [["A", "B"], [1, 2]]


# ---------------------------------------------------------------- Excel
def test_excel_write_professional_formatting():
    call("excel_write", {"path": "rel.xlsx", "sheets": {"Vendas": VENDAS},
                         "formats": {"Vendas": {"Valor": "money"}}, "totals": {"Vendas": ["Valor", "Qtd"]},
                         "conditional": {"Vendas": {"Valor": "bar"}},
                         "charts": [{"sheet": "Vendas", "type": "bar", "category_col": "Vendedor", "value_cols": ["Valor"]}]})
    ws = load_workbook(config.workspace() / "rel.xlsx")["Vendas"]
    assert ws["E2"].number_format == "R$ #,##0.00"
    assert ws["A8"].value == "TOTAL" and ws["E8"].value == "=SUM(E2:E7)"
    assert len(ws._charts) == 1 and len(ws.conditional_formatting) == 1


def test_excel_dashboard_native_with_live_formulas():
    src = _write_vendas()
    r = call("excel_dashboard", {"path": "dash.xlsx", "source": src, "title": "Vendas 2026",
                                 "kpis": [{"label": "Faturamento", "column": "Valor", "agg": "sum", "format": "money"},
                                          {"label": "Pedidos", "column": "Valor", "agg": "count"},
                                          {"label": "Vendedores", "column": "Vendedor", "agg": "nunique"}],
                                 "charts": [{"type": "bar", "x": "Região", "y": "Valor"},
                                            {"type": "line", "x": "Data", "y": "Valor", "period": "month"},
                                            {"type": "pie", "x": "Produto"}]})
    assert r["sheets"] == ["Dashboard", "Dados", "Resumos"]
    wb = load_workbook(config.workspace() / "dash.xlsx")
    d = wb["Dashboard"]
    assert d["B2"].value == "Vendas 2026" and len(d._charts) == 3
    formulas = [c.value for row in d.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert "=SUM(Dados!$E:$E)" in formulas and "=COUNTA(Dados!$E:$E)-1" in formulas
    assert 3 in [c.value for row in d.iter_rows() for c in row]  # distintos
    res = wb["Resumos"]
    assert res["B2"].value.startswith("=SUMIF(Dados!$B:$B,Resumos!$A$2") and res["A2"].value == "Sul"
    assert [res.cell(r, 4).value for r in (2, 3, 4)] == ["2026-01", "2026-02", "2026-03"]
    assert wb.calculation.fullCalcOnLoad


def test_excel_update_existing_file_keeps_data():
    src = _write_vendas()
    call("excel_update", {"path": src, "sheet": "Vendas", "cells": {"H1": "Obs", "H2": "=E2*2"},
                          "append_rows": [["01/04/2026", "Sul", "Ana", "B", 10, 1]]})
    ws = load_workbook(config.workspace() / src)["Vendas"]
    assert ws["H2"].value == "=E2*2" and ws.max_row == 8 and ws["A1"].font.bold


def test_dashboard_data_interactive_html(headless):
    src = _write_vendas()
    call("dashboard_data", {"path": "painel.html", "source": src, "title": "Painel de Vendas",
                            "kpis": [{"label": "Total", "column": "Valor", "agg": "sum", "format": "money"},
                                     {"label": "Pedidos", "column": "Valor", "agg": "count", "format": "int"}],
                            "charts": [{"type": "bar", "x": "Vendedor", "y": "Valor"},
                                       {"type": "line", "x": "Data", "y": "Valor", "period": "month"}],
                            "filters": ["Região"], "date_column": "Data"})
    from playwright.sync_api import sync_playwright
    import os
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path=os.environ.get("EMANUSIA_CHROMIUM"), headless=True)
        pg = b.new_page()
        errs = []
        pg.on("pageerror", lambda e: errs.append(str(e)))
        pg.goto((config.workspace() / "painel.html").as_uri())
        total = pg.inner_text("#k0")
        pg.select_option("select", "Norte")
        norte = pg.inner_text("#k0"), pg.inner_text("#k1"), pg.inner_text("#cnt")
        pg.select_option("select", "")
        pg.fill("#d_from", "2026-03-01")
        marco = pg.inner_text("#k1")
        b.close()
    assert not errs
    assert "7.250,50" in total and "1.750,50" in norte[0] and norte[1] == "2" and norte[2].startswith("2 de 6")
    assert marco == "2"


def test_dashboard_bar_axes_are_categories_not_dates(headless):
    """Regressão: o eixo de data de um gráfico mensal contaminava os gráficos de barra."""
    call("dashboard_data", {"path": "p2.html", "source": VENDAS, "title": "t",
                            "charts": [{"type": "line", "x": "Data", "y": "Valor", "period": "month"},
                                       {"type": "bar", "x": "Vendedor", "y": "Valor"}], "date_column": "Data"})
    from playwright.sync_api import sync_playwright
    import os
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path=os.environ.get("EMANUSIA_CHROMIUM"), headless=True)
        pg = b.new_page()
        pg.goto((config.workspace() / "p2.html").as_uri())
        types = pg.evaluate("[document.getElementById('c0').layout.xaxis.type, document.getElementById('c1').layout.xaxis.type]")
        ticks = pg.inner_text("#c1")
        b.close()
    assert types == ["category", "category"] and "Ana" in ticks and "Rui" in ticks
